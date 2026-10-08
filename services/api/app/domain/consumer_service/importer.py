"""Official L'Oreal business workbook importer.

The importer keeps two copies of source facts:
1. the original workbook snapshot on disk;
2. the original row values in ``service_source_record``.

Normalized rows are separate read-only facts. AI observations and mutable demo
conversation state must not be written into these tables.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.domain.consumer_service.mapping import (
    CONTENT_TYPE_MAP,
    DATASET_ID,
    MAPPING_VERSION,
    ROLE_MAP,
    SHEETS,
    STATUS_MAP,
)
from app.domain.consumer_service.models import (
    AdverseReactionRow,
    DatasetRow,
    DetailMixin,
    ImportBatchRow,
    LogisticsRow,
    OfflinePaymentRow,
    ReshipExchangeRow,
    ReturnRefundRow,
    ServiceConversationRow,
    ServiceEventRow,
    ServiceMessageRow,
    ServiceOrderRow,
    SourceRecordRow,
    WorkOrderRow,
)


class ImportValidationError(ValueError):
    """The workbook cannot be safely published as a business-data batch."""

    def __init__(self, report: dict[str, Any]):
        self.report = report
        super().__init__("业务数据质量检查未通过，批次未发布")


@dataclass(frozen=True)
class ParsedRow:
    sheet_name: str
    row_number: int
    raw_values: dict[str, Any]
    cell_metadata: dict[str, dict[str, Any]]
    normalized: dict[str, Any]
    source_record_id: str


@dataclass(frozen=True)
class ImportResult:
    batch_id: str
    source_sha256: str
    snapshot_path: str
    report_path: str
    reused: bool
    counts: dict[str, int]
    report: dict[str, Any]


def _jsonable(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, bytes):
        return value.hex()
    return value


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    return value


def _parse_datetime(value: Any) -> datetime | None:
    value = _clean(value)
    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, str):
        candidate = value.replace("T", " ")
        prefix = re.match(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}", candidate)
        if prefix:
            candidate = prefix.group(0)
        try:
            return datetime.fromisoformat(candidate)
        except ValueError:
            return None
    return None


def _parse_int(value: Any) -> int | None:
    value = _clean(value)
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _minor_units(value: Any) -> int | None:
    value = _clean(value)
    if value is None:
        return None
    try:
        return int(
            (Decimal(str(value)) * Decimal("100")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        )
    except Exception:
        return None


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _record_id(batch_id: str, sheet_name: str, row_number: int) -> str:
    value = uuid.uuid5(
        uuid.NAMESPACE_URL, f"{DATASET_ID}:{MAPPING_VERSION}:{batch_id}:{sheet_name}:{row_number}"
    )
    return f"src-{value.hex}"


def _fact_id(batch_id: str, kind: str, key: str) -> str:
    value = uuid.uuid5(
        uuid.NAMESPACE_URL, f"{DATASET_ID}:{MAPPING_VERSION}:{batch_id}:{kind}:{key}"
    )
    return f"svc-{value.hex}"


def _event_id(source_record_id: str, event_type: str) -> str:
    value = uuid.uuid5(uuid.NAMESPACE_URL, f"{source_record_id}:{event_type}")
    return f"evt-{value.hex}"


def _load_rows(source_path: Path, batch_id: str) -> dict[str, list[ParsedRow]]:
    workbook = load_workbook(source_path, read_only=True, data_only=False)
    try:
        actual = set(workbook.sheetnames)
        expected = set(SHEETS)
        if missing := sorted(expected - actual):
            raise ValueError(f"缺少数据表：{','.join(missing)}")

        parsed: dict[str, list[ParsedRow]] = {}
        for sheet_name, mapping in SHEETS.items():
            worksheet = workbook[sheet_name]
            rows = worksheet.iter_rows()
            try:
                header_cells = next(rows)
            except StopIteration as exc:
                raise ValueError(f"数据表为空：{sheet_name}") from exc

            headers = [
                str(cell.value).strip() if cell.value is not None else "" for cell in header_cells
            ]
            if len(headers) != len(set(headers)) or "" in headers:
                raise ValueError(f"表头无效：{sheet_name}")

            expected_headers = set(mapping.fields)
            actual_headers = set(headers)
            if missing_headers := sorted(expected_headers - actual_headers):
                raise ValueError(f"{sheet_name} 缺少字段：{','.join(missing_headers)}")
            if extra_headers := sorted(actual_headers - expected_headers):
                raise ValueError(f"{sheet_name} 存在未映射字段：{','.join(extra_headers)}")

            parsed_rows: list[ParsedRow] = []
            for row_number, cells in enumerate(rows, start=2):
                raw_values = {
                    header: _jsonable(cell.value)
                    for header, cell in zip(headers, cells, strict=True)
                }
                if all(value is None or value == "" for value in raw_values.values()):
                    continue
                cell_metadata = {
                    header: {
                        "coordinate": getattr(
                            cell, "coordinate", f"{get_column_letter(column_number)}{row_number}"
                        ),
                        "dataType": getattr(cell, "data_type", "n"),
                        "numberFormat": getattr(cell, "number_format", "General"),
                    }
                    for column_number, (header, cell) in enumerate(
                        zip(headers, cells, strict=True), start=1
                    )
                }
                normalized = {
                    target: _clean(raw_values[source]) for source, target in mapping.fields.items()
                }
                parsed_rows.append(
                    ParsedRow(
                        sheet_name=sheet_name,
                        row_number=row_number,
                        raw_values=raw_values,
                        cell_metadata=cell_metadata,
                        normalized=normalized,
                        source_record_id=_record_id(batch_id, sheet_name, row_number),
                    )
                )
            parsed[sheet_name] = parsed_rows
        return parsed
    finally:
        workbook.close()


def _issue(
    issues: list[dict[str, Any]],
    *,
    code: str,
    severity: str,
    sheet: str | None = None,
    count: int = 1,
    samples: list[Any] | None = None,
) -> None:
    item: dict[str, Any] = {"code": code, "severity": severity, "count": count}
    if sheet is not None:
        item["sheet"] = sheet
    if samples:
        item["samples"] = samples[:5]
    issues.append(item)


def _quality_report(
    source_path: Path,
    source_sha256: str,
    parsed: dict[str, list[ParsedRow]],
    batch_id: str,
) -> dict[str, Any]:
    issues: list[dict[str, Any]] = []
    counts = {sheet: len(rows) for sheet, rows in parsed.items()}
    duplicate_keys: dict[str, list[str]] = {}
    for sheet_name, rows in parsed.items():
        mapping = SHEETS[sheet_name]
        key_values = [row.normalized.get(mapping.key) for row in rows]
        missing_count = sum(value is None for value in key_values)
        if missing_count:
            _issue(
                issues,
                code="missing_business_key",
                severity="blocking",
                sheet=sheet_name,
                count=missing_count,
            )
        keys = [str(value) for value in key_values if value is not None]
        duplicates = sorted(key for key, count in Counter(keys).items() if count > 1)
        if duplicates:
            duplicate_keys[sheet_name] = duplicates
            _issue(
                issues,
                code="duplicate_business_key",
                severity="blocking",
                sheet=sheet_name,
                count=sum(Counter(keys)[key] - 1 for key in duplicates),
                samples=duplicates,
            )

        for field in mapping.datetime_fields:
            invalid = [
                row.row_number
                for row in rows
                if row.normalized.get(field) is not None
                and _parse_datetime(row.normalized[field]) is None
            ]
            if invalid:
                _issue(
                    issues,
                    code="invalid_datetime",
                    severity="blocking",
                    sheet=sheet_name,
                    count=len(invalid),
                    samples=invalid,
                )
        for field in mapping.money_fields:
            invalid = [
                row.row_number
                for row in rows
                if row.normalized.get(field) is not None
                and _minor_units(row.normalized[field]) is None
            ]
            if invalid:
                _issue(
                    issues,
                    code="invalid_money",
                    severity="blocking",
                    sheet=sheet_name,
                    count=len(invalid),
                    samples=invalid,
                )
        for field in mapping.integer_fields:
            invalid = [
                row.row_number
                for row in rows
                if row.normalized.get(field) is not None
                and _parse_int(row.normalized[field]) is None
            ]
            if invalid:
                _issue(
                    issues,
                    code="invalid_integer",
                    severity="blocking",
                    sheet=sheet_name,
                    count=len(invalid),
                    samples=invalid,
                )

    chat_rows = parsed["聊天记录"]
    chat_conversations = {
        row.normalized["conversation_id"]
        for row in chat_rows
        if row.normalized.get("conversation_id") is not None
    }
    order_rows = parsed["订单"]
    order_ids = {
        row.normalized["order_id"]
        for row in order_rows
        if row.normalized.get("order_id") is not None
    }
    work_rows = [row for name in SHEETS if name not in {"聊天记录", "订单"} for row in parsed[name]]
    work_ids = {
        row.normalized["work_order_id"]
        for row in work_rows
        if row.normalized.get("work_order_id") is not None
    }

    def relation_check(
        rows: list[ParsedRow],
        field: str,
        values: set[Any],
        code: str,
        sheet: str,
    ) -> None:
        unmatched = [
            row.normalized.get(field)
            for row in rows
            if row.normalized.get(field) is not None and row.normalized.get(field) not in values
        ]
        if unmatched:
            _issue(
                issues,
                code=code,
                severity="blocking",
                sheet=sheet,
                count=len(unmatched),
                samples=sorted({str(value) for value in unmatched}),
            )

    relation_check(
        order_rows, "conversation_id", chat_conversations, "order_conversation_unmatched", "订单"
    )
    for sheet_name in SHEETS:
        if sheet_name in {"聊天记录", "订单"}:
            continue
        rows = parsed[sheet_name]
        relation_check(
            rows,
            "conversation_id",
            chat_conversations,
            "work_order_conversation_unmatched",
            sheet_name,
        )
        relation_check(rows, "order_id", order_ids, "work_order_order_unmatched", sheet_name)
    relation_check(chat_rows, "order_id", order_ids, "message_order_unmatched", "聊天记录")
    relation_check(chat_rows, "work_order_id", work_ids, "message_work_order_unmatched", "聊天记录")

    def conflicts(
        rows: list[ParsedRow], group_field: str, value_field: str
    ) -> dict[str, list[str]]:
        grouped: dict[str, set[str]] = defaultdict(set)
        for row in rows:
            group = row.normalized.get(group_field)
            value = row.normalized.get(value_field)
            if group is not None and value is not None:
                grouped[str(group)].add(str(value))
        return {group: sorted(values) for group, values in grouped.items() if len(values) > 1}

    order_conflicts = conflicts(order_rows, "conversation_id", "order_id")
    work_conflicts = conflicts(work_rows, "conversation_id", "work_order_id")
    if order_conflicts:
        _issue(
            issues,
            code="conversation_multiple_orders",
            severity="blocking",
            count=len(order_conflicts),
            samples=list(order_conflicts),
        )
    if work_conflicts:
        _issue(
            issues,
            code="conversation_multiple_work_orders",
            severity="blocking",
            count=len(work_conflicts),
            samples=list(work_conflicts),
        )

    alias_by_conversation: dict[str, set[str]] = defaultdict(set)
    for rows in parsed.values():
        for row in rows:
            conversation_id = row.normalized.get("conversation_id")
            buyer_alias = row.normalized.get("buyer_alias")
            if conversation_id is not None and buyer_alias is not None:
                alias_by_conversation[str(conversation_id)].add(str(buyer_alias))
    alias_conflicts = {
        conversation: sorted(values)
        for conversation, values in alias_by_conversation.items()
        if len(values) > 1
    }
    if alias_conflicts:
        _issue(
            issues,
            code="conversation_alias_conflict",
            severity="warning",
            count=len(alias_conflicts),
            samples=list(alias_conflicts),
        )

    unknown_roles = [
        row.row_number for row in chat_rows if row.normalized.get("source_role") not in ROLE_MAP
    ]
    if unknown_roles:
        _issue(
            issues,
            code="unknown_message_role",
            severity="blocking",
            sheet="聊天记录",
            count=len(unknown_roles),
            samples=unknown_roles,
        )
    unknown_content_types = [
        row.row_number
        for row in chat_rows
        if row.normalized.get("source_content_type") not in CONTENT_TYPE_MAP
    ]
    if unknown_content_types:
        _issue(
            issues,
            code="unknown_content_type",
            severity="blocking",
            sheet="聊天记录",
            count=len(unknown_content_types),
            samples=unknown_content_types,
        )

    image_rows = [row for row in chat_rows if row.normalized.get("source_content_type") == "图片"]
    status_counts: dict[str, dict[str, int]] = {}
    for sheet_name, mapping in SHEETS.items():
        if mapping.status_field is None:
            continue
        status_counts[sheet_name] = dict(
            Counter(
                str(row.normalized.get(mapping.status_field))
                for row in parsed[sheet_name]
                if row.normalized.get(mapping.status_field) is not None
            )
        )

    blocking_issues = [issue for issue in issues if issue["severity"] == "blocking"]
    return {
        "datasetId": DATASET_ID,
        "mappingVersion": MAPPING_VERSION,
        "source": {
            "name": source_path.name,
            "sha256": source_sha256,
            "sourceKind": "official_fictional_mock",
        },
        "batchId": batch_id,
        "qualityStatus": "failed" if blocking_issues else "passed",
        "counts": {
            "sheets": len(parsed),
            "messages": len(chat_rows),
            "orders": len(order_rows),
            "workOrders": len(work_rows),
            "conversations": len(chat_conversations),
            "imageMessages": len(image_rows),
        },
        "sheetCounts": counts,
        "rawFieldMapping": {sheet_name: mapping.fields for sheet_name, mapping in SHEETS.items()},
        "duplicateKeys": duplicate_keys,
        "relationships": {
            "orderConversationUnmatched": sum(
                issue["count"]
                for issue in issues
                if issue["code"] == "order_conversation_unmatched"
            ),
            "workOrderConversationUnmatched": sum(
                issue["count"]
                for issue in issues
                if issue["code"] == "work_order_conversation_unmatched"
            ),
            "workOrderOrderUnmatched": sum(
                issue["count"] for issue in issues if issue["code"] == "work_order_order_unmatched"
            ),
            "messageOrderUnmatched": sum(
                issue["count"] for issue in issues if issue["code"] == "message_order_unmatched"
            ),
            "messageWorkOrderUnmatched": sum(
                issue["count"]
                for issue in issues
                if issue["code"] == "message_work_order_unmatched"
            ),
            "conversationOrderConflicts": order_conflicts,
            "conversationWorkOrderConflicts": work_conflicts,
            "conversationAliasConflicts": alias_conflicts,
        },
        "statusCounts": status_counts,
        "imageReferences": {
            "count": len(image_rows),
            "withPath": sum(1 for row in image_rows if row.normalized.get("image_ref")),
        },
        "issues": issues,
        "blockingIssues": blocking_issues,
        "warnings": [issue for issue in issues if issue["severity"] == "warning"],
    }


def _normal_values(row: ParsedRow) -> dict[str, Any]:
    return {key: _jsonable(value) for key, value in row.normalized.items()}


def _create_conversation_rows(
    batch_id: str, parsed: dict[str, list[ParsedRow]]
) -> tuple[dict[str, str], list[ServiceConversationRow]]:
    grouped: dict[str, dict[str, Any]] = defaultdict(lambda: {"aliases": set(), "timestamps": []})
    for rows in parsed.values():
        for row in rows:
            conversation_id = row.normalized.get("conversation_id")
            if conversation_id is None:
                continue
            item = grouped[str(conversation_id)]
            alias = row.normalized.get("buyer_alias")
            if alias is not None:
                item["aliases"].add(str(alias))
            for field in (
                "sent_at",
                "ordered_at",
                "paid_at",
                "shipped_at",
                "created_at",
                "completed_at",
            ):
                value = row.normalized.get(field)
                parsed_time = _parse_datetime(value)
                if parsed_time is not None:
                    item["timestamps"].append(parsed_time)

    records: dict[str, str] = {}
    rows: list[ServiceConversationRow] = []
    for conversation_id, item in grouped.items():
        record_id = _fact_id(batch_id, "conversation", conversation_id)
        aliases = sorted(item["aliases"])
        timestamps = item["timestamps"]
        rows.append(
            ServiceConversationRow(
                record_id=record_id,
                batch_id=batch_id,
                conversation_id=conversation_id,
                buyer_alias=aliases[0] if aliases else "",
                aliases_json=_json(aliases),
                first_at=min(timestamps) if timestamps else None,
                last_at=max(timestamps) if timestamps else None,
            )
        )
        records[conversation_id] = record_id
    return records, rows


def _create_source_rows(batch_id: str, parsed: dict[str, list[ParsedRow]]) -> list[SourceRecordRow]:
    return [
        SourceRecordRow(
            record_id=row.source_record_id,
            batch_id=batch_id,
            sheet_name=row.sheet_name,
            row_number=row.row_number,
            business_key=(
                str(row.normalized.get(SHEETS[row.sheet_name].key))
                if row.normalized.get(SHEETS[row.sheet_name].key) is not None
                else None
            ),
            values_json=_json(row.raw_values),
            cell_metadata_json=_json(row.cell_metadata),
            row_sha256=hashlib.sha256(_json(row.raw_values).encode("utf-8")).hexdigest(),
        )
        for rows in parsed.values()
        for row in rows
    ]


def _create_order_rows(
    batch_id: str,
    rows: list[ParsedRow],
    conversations: dict[str, str],
) -> tuple[dict[str, str], list[ServiceOrderRow]]:
    records: dict[str, str] = {}
    result: list[ServiceOrderRow] = []
    for row in rows:
        data = row.normalized
        order_id = str(data["order_id"])
        record_id = _fact_id(batch_id, "order", order_id)
        records[order_id] = record_id
        result.append(
            ServiceOrderRow(
                record_id=record_id,
                batch_id=batch_id,
                source_record_id=row.source_record_id,
                conversation_record_id=conversations[str(data["conversation_id"])],
                conversation_id=str(data["conversation_id"]),
                buyer_alias=str(data.get("buyer_alias") or ""),
                shop=data.get("shop"),
                order_id=order_id,
                sku=data.get("sku"),
                product_name=data.get("product_name"),
                quantity=_parse_int(data.get("quantity")),
                unit_price_minor=_minor_units(data.get("unit_price")),
                paid_amount_minor=_minor_units(data.get("paid_amount")),
                source_status=data.get("source_status"),
                ordered_at=_parse_datetime(data.get("ordered_at")) or datetime.min,
                paid_at=_parse_datetime(data.get("paid_at")),
                shipped_at=_parse_datetime(data.get("shipped_at")),
                carrier=data.get("carrier"),
                tracking_no=data.get("tracking_no"),
                province=data.get("province"),
                city=data.get("city"),
                gift=data.get("gift"),
                buyer_note=data.get("buyer_note"),
            )
        )
    return records, result


def _create_work_order_rows(
    batch_id: str,
    parsed: dict[str, list[ParsedRow]],
    conversations: dict[str, str],
    orders: dict[str, str],
) -> tuple[dict[str, str], list[WorkOrderRow], list[DetailMixin]]:
    records: dict[str, str] = {}
    result: list[WorkOrderRow] = []
    details: list[DetailMixin] = []
    for sheet_name in SHEETS:
        if sheet_name in {"聊天记录", "订单"}:
            continue
        mapping = SHEETS[sheet_name]
        for row in parsed[sheet_name]:
            data = row.normalized
            work_order_id = str(data["work_order_id"])
            record_id = _fact_id(batch_id, mapping.entity, work_order_id)
            records[work_order_id] = record_id
            order_id = data.get("order_id")
            detail_type: type[DetailMixin]
            if mapping.entity == "reship_exchange":
                detail_type = ReshipExchangeRow
                detail_kwargs = {
                    "service_type": data.get("service_type"),
                    "reason": data.get("reason"),
                    "original_tracking_no": data.get("original_tracking_no"),
                    "reship_tracking_no": data.get("reship_tracking_no"),
                }
            elif mapping.entity == "offline_payment":
                detail_type = OfflinePaymentRow
                detail_kwargs = {
                    "refund_amount_minor": _minor_units(data.get("refund_amount")),
                    "transfer_status": data.get("transfer_status"),
                    "refund_reason": data.get("refund_reason"),
                }
            elif mapping.entity == "logistics":
                detail_type = LogisticsRow
                detail_kwargs = {
                    "issue_type": data.get("issue_type"),
                    "solution": data.get("solution"),
                    "tracking_no": data.get("tracking_no"),
                }
            elif mapping.entity == "adverse_reaction":
                detail_type = AdverseReactionRow
                detail_kwargs = {
                    "symptom_description": data.get("symptom_description"),
                    "sought_medical_care": data.get("sought_medical_care"),
                    "stopped_use": data.get("stopped_use"),
                }
            else:
                detail_type = ReturnRefundRow
                detail_kwargs = {
                    "return_reason": data.get("return_reason"),
                    "refund_id": data.get("refund_id"),
                    "abnormal": data.get("abnormal"),
                }
            result.append(
                WorkOrderRow(
                    record_id=record_id,
                    batch_id=batch_id,
                    source_record_id=row.source_record_id,
                    conversation_record_id=conversations[str(data["conversation_id"])],
                    conversation_id=str(data["conversation_id"]),
                    buyer_alias=str(data.get("buyer_alias") or ""),
                    shop=data.get("shop"),
                    work_order_id=work_order_id,
                    work_order_type=mapping.entity,
                    order_id=str(order_id) if order_id is not None else None,
                    order_record_id=orders.get(str(order_id)) if order_id is not None else None,
                    source_status=str(data.get("source_status") or ""),
                    normalized_status=STATUS_MAP.get(str(data.get("source_status")), "unknown"),
                    handler=data.get("handler"),
                    created_at=_parse_datetime(data.get("created_at")) or datetime.min,
                    completed_at=_parse_datetime(data.get("completed_at")),
                )
            )
            details.append(
                detail_type(
                    work_order_record_id=record_id,
                    fields_json=_json(_normal_values(row)),
                    **detail_kwargs,
                )
            )
    return records, result, details


def _create_message_rows(
    batch_id: str,
    rows: list[ParsedRow],
    conversations: dict[str, str],
    orders: dict[str, str],
    work_orders: dict[str, str],
) -> list[ServiceMessageRow]:
    result: list[ServiceMessageRow] = []
    for row in rows:
        data = row.normalized
        conversation_id = str(data["conversation_id"])
        order_id = data.get("order_id")
        work_order_id = data.get("work_order_id")
        image_ref = data.get("image_ref")
        result.append(
            ServiceMessageRow(
                record_id=_fact_id(batch_id, "message", str(data["message_id"])),
                batch_id=batch_id,
                source_record_id=row.source_record_id,
                conversation_record_id=conversations[conversation_id],
                conversation_id=conversation_id,
                buyer_alias=str(data.get("buyer_alias") or ""),
                shop=data.get("shop"),
                message_id=str(data["message_id"]),
                sequence=_parse_int(data.get("sequence")) or 0,
                sent_at=_parse_datetime(data.get("sent_at")) or datetime.min,
                source_role=str(data.get("source_role") or ""),
                sender_role=ROLE_MAP.get(str(data.get("source_role")), "unknown"),
                sender=data.get("sender"),
                body=data.get("body"),
                scene_major=data.get("scene_major"),
                scene_minor=data.get("scene_minor"),
                is_target_buyer_message=_parse_int(data.get("is_target_buyer_message")),
                source_content_type=str(data.get("source_content_type") or ""),
                content_type=CONTENT_TYPE_MAP.get(str(data.get("source_content_type")), "unknown"),
                chat_content=data.get("chat_content"),
                category=data.get("category"),
                image_ref=image_ref,
                image_state="reference" if image_ref else "not_applicable",
                order_id=str(order_id) if order_id is not None else None,
                work_order_id=str(work_order_id) if work_order_id is not None else None,
                order_record_id=orders.get(str(order_id)) if order_id is not None else None,
                work_order_record_id=(
                    work_orders.get(str(work_order_id)) if work_order_id is not None else None
                ),
            )
        )
    return result


def _create_events(
    batch_id: str,
    parsed: dict[str, list[ParsedRow]],
    conversations: dict[str, str],
    orders: dict[str, str],
    work_orders: dict[str, str],
) -> list[ServiceEventRow]:
    events: list[ServiceEventRow] = []
    for row in parsed["聊天记录"]:
        data = row.normalized
        conversation_id = str(data["conversation_id"])
        events.append(
            ServiceEventRow(
                event_id=_event_id(row.source_record_id, "message"),
                batch_id=batch_id,
                conversation_record_id=conversations[conversation_id],
                conversation_id=conversation_id,
                source_record_id=row.source_record_id,
                entity_record_id=_fact_id(batch_id, "message", str(data["message_id"])),
                entity_type="message",
                event_type="message",
                occurred_at=_parse_datetime(data.get("sent_at")) or datetime.min,
            )
        )
    for row in parsed["订单"]:
        data = row.normalized
        conversation_id = str(data["conversation_id"])
        order_id = str(data["order_id"])
        events.append(
            ServiceEventRow(
                event_id=_event_id(row.source_record_id, "order"),
                batch_id=batch_id,
                conversation_record_id=conversations[conversation_id],
                conversation_id=conversation_id,
                source_record_id=row.source_record_id,
                entity_record_id=orders[order_id],
                entity_type="order",
                event_type="order",
                occurred_at=_parse_datetime(data.get("ordered_at")) or datetime.min,
            )
        )
    for sheet_name in SHEETS:
        if sheet_name in {"聊天记录", "订单"}:
            continue
        mapping = SHEETS[sheet_name]
        for row in parsed[sheet_name]:
            data = row.normalized
            conversation_id = str(data["conversation_id"])
            work_order_id = str(data["work_order_id"])
            events.append(
                ServiceEventRow(
                    event_id=_event_id(row.source_record_id, "work_order"),
                    batch_id=batch_id,
                    conversation_record_id=conversations[conversation_id],
                    conversation_id=conversation_id,
                    source_record_id=row.source_record_id,
                    entity_record_id=work_orders[work_order_id],
                    entity_type="work_order",
                    event_type=f"work_order:{mapping.entity}",
                    occurred_at=_parse_datetime(data.get("created_at")) or datetime.min,
                )
            )
    return events


def _write_report(report_path: Path, report: dict[str, Any]) -> None:
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str) + "\n",
        encoding="utf-8",
    )


def import_business_workbook(
    session: Session,
    source_path: Path,
    *,
    repo_root: Path | None = None,
) -> ImportResult:
    source_path = source_path.resolve()
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    repo_root = repo_root or Path(__file__).resolve().parents[5]
    source_sha256 = _sha256(source_path)
    batch_id = f"loreal-{source_sha256[:24]}"

    snapshot_path = (
        repo_root / "data" / "source" / "loreal-official-mock" / (f"source-{source_sha256}.xlsx")
    )
    snapshot_path.parent.mkdir(parents=True, exist_ok=True)
    if not snapshot_path.exists():
        shutil.copy2(source_path, snapshot_path)

    report_path = repo_root / "data" / "normalized" / "loreal" / batch_id / "data-quality.json"
    existing = session.scalar(
        select(ImportBatchRow).where(
            ImportBatchRow.source_sha256 == source_sha256,
            ImportBatchRow.mapping_version == MAPPING_VERSION,
        )
    )
    if existing is not None:
        report = json.loads(existing.report_json)
        _write_report(report_path, report)
        return ImportResult(
            batch_id=existing.batch_id,
            source_sha256=source_sha256,
            snapshot_path=existing.snapshot_path,
            report_path=str(report_path),
            reused=True,
            counts=report["counts"],
            report=report,
        )

    parsed = _load_rows(source_path, batch_id)
    report = _quality_report(source_path, source_sha256, parsed, batch_id)
    _write_report(report_path, report)
    if report["blockingIssues"]:
        raise ImportValidationError(report)

    conversation_records, conversation_rows = _create_conversation_rows(batch_id, parsed)
    order_records, order_rows = _create_order_rows(batch_id, parsed["订单"], conversation_records)
    work_order_records, work_order_rows, detail_rows = _create_work_order_rows(
        batch_id, parsed, conversation_records, order_records
    )
    message_rows = _create_message_rows(
        batch_id,
        parsed["聊天记录"],
        conversation_records,
        order_records,
        work_order_records,
    )
    event_rows = _create_events(
        batch_id, parsed, conversation_records, order_records, work_order_records
    )
    source_rows = _create_source_rows(batch_id, parsed)

    batch = ImportBatchRow(
        batch_id=batch_id,
        source_sha256=source_sha256,
        mapping_version=MAPPING_VERSION,
        source_name=source_path.name,
        snapshot_path=str(snapshot_path),
        source_kind="official_fictional_mock",
        status="published",
        report_json=_json(report),
    )
    session.add(batch)
    session.add_all(source_rows)
    session.flush()
    session.add_all(conversation_rows)
    session.flush()
    session.add_all(order_rows)
    session.flush()
    session.add_all(work_order_rows)
    session.flush()
    session.add_all(detail_rows)
    session.flush()
    session.add_all(message_rows)
    session.add_all(event_rows)

    dataset = session.get(DatasetRow, DATASET_ID)
    if dataset is None:
        session.add(DatasetRow(dataset_id=DATASET_ID, active_batch_id=batch_id))
    else:
        dataset.active_batch_id = batch_id
    session.flush()
    from app.domain.consumer_service.risk import seed_rule_alerts

    seed_rule_alerts(session, batch_id)

    return ImportResult(
        batch_id=batch_id,
        source_sha256=source_sha256,
        snapshot_path=str(snapshot_path),
        report_path=str(report_path),
        reused=False,
        counts=report["counts"],
        report=report,
    )
